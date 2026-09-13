#!/usr/bin/env python3
"""Deterministic helpers for the tailored-resume-agent Codex Skill."""

from __future__ import annotations

import argparse
import hashlib
import json
import logging
import os
import re
import secrets
import shutil
import subprocess
import sys
import tempfile
from datetime import datetime
from pathlib import Path
from typing import Any
from urllib.parse import urlparse


SCHEMA_VERSION = 1
HTML_SUFFIXES = {".html", ".htm"}
DOCX_SUFFIXES = {".docx"}
EXTRACTABLE_SUFFIXES = {".html", ".htm", ".docx", ".pdf", ".txt", ".md"}
PLACEHOLDER_RE = re.compile(r"\[(?:TODO|待确认|待补|未确认|TBD)[^\]]*\]", re.IGNORECASE)
NUMBER_RE = re.compile(r"\d+(?:[.,]\d+)?%?")
STRONG_CLAIM_RE = re.compile(r"主导|统筹|带领|独立负责|全权负责")
VAGUE_CLAIM_RE = re.compile(
    r"沟通能力强|学习能力强|执行力强|责任心强|抗压能力强|认真负责|积极主动|团队协作能力强"
)
RESULT_SIGNAL_RE = re.compile(
    r"规模|频率|周期|范围|排名|对比|覆盖|节省|提升|下降|降低|增长|上线|达成|完成|从.+到"
)
INVALID_FILENAME_CHARS = re.compile(r"[/\\:*?\"<>|\x00-\x1f]")
UNRESOLVED_METADATA = {"", "待识别", "待确认", "unknown", "TBD"}


class AgentError(RuntimeError):
    """A user-actionable workflow error."""


def now_iso() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")


def read_json(path: Path) -> Any:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except FileNotFoundError as exc:
        raise AgentError(f"文件不存在：{path}") from exc
    except json.JSONDecodeError as exc:
        raise AgentError(f"JSON 格式错误：{path}: {exc}") from exc


def write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


def safe_filename_stem(value: str) -> str:
    stem = INVALID_FILENAME_CHARS.sub("-", value.strip()).strip(". ")
    stem = re.sub(r"-{2,}", "-", stem)
    if not stem:
        raise AgentError("文件名不能为空。")
    return stem


def slugify(value: str, limit: int = 36) -> str:
    value = INVALID_FILENAME_CHARS.sub("-", value.strip().lower())
    value = re.sub(r"\s+", "-", value)
    value = re.sub(r"-{2,}", "-", value).strip("-")
    return value[:limit].strip("-") or "resume"


def unique_destination(directory: Path, filename: str) -> Path:
    candidate = directory / filename
    if not candidate.exists():
        return candidate
    stem = candidate.stem
    suffix = candidate.suffix
    counter = 2
    while True:
        candidate = directory / f"{stem}-{counter}{suffix}"
        if not candidate.exists():
            return candidate
        counter += 1


def is_http_url(value: str) -> bool:
    parsed = urlparse(value)
    return parsed.scheme in {"http", "https"} and bool(parsed.netloc)


def extract_text(path: Path) -> str:
    if not path.exists():
        raise AgentError(f"输入文件不存在：{path}")
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        return path.read_text(encoding="utf-8", errors="replace")
    if suffix in HTML_SUFFIXES:
        try:
            from lxml import html
        except ImportError as exc:
            raise AgentError("读取 HTML 需要 lxml，请使用 Codex workspace dependency 的 Python。") from exc
        document = html.parse(str(path))
        for node in document.xpath("//style|//script|//noscript"):
            node.drop_tree()
        return " ".join(document.getroot().text_content().split())
    if suffix in DOCX_SUFFIXES:
        try:
            from docx import Document
        except ImportError as exc:
            raise AgentError("读取 DOCX 需要 python-docx，请使用 Codex workspace dependency 的 Python。") from exc
        document = Document(str(path))
        chunks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        chunks.append(cell.text)
        return "\n".join(chunks)
    if suffix == ".pdf":
        try:
            import pdfplumber
        except ImportError as exc:
            raise AgentError("读取 PDF 需要 pdfplumber，请使用 Codex workspace dependency 的 Python。") from exc
        chunks: list[str] = []
        with pdfplumber.open(str(path)) as pdf:
            for page in pdf.pages:
                chunks.append(page.extract_text() or "")
        return "\n".join(chunks)
    raise AgentError(f"不支持的文件类型：{path.suffix}")


def run_dir_base(store: Path) -> Path:
    return store / "runs"


def ensure_store(store: Path) -> None:
    profile = store / "profile.json"
    evidence = store / "evidence.json"
    if not profile.exists() or not evidence.exists():
        raise AgentError(
            f"事实库尚未初始化：{store}\n"
            "请先运行 resume_agent.py init --store <store> --resume <resume>。"
        )


def init_store(args: argparse.Namespace) -> dict[str, Any]:
    store = Path(args.store).expanduser().resolve()
    resume = Path(args.resume).expanduser().resolve()
    if resume.suffix.lower() not in EXTRACTABLE_SUFFIXES:
        raise AgentError(f"不支持作为简历输入的文件类型：{resume.suffix}")

    store.mkdir(parents=True, exist_ok=True)
    profile_path = store / "profile.json"
    evidence_path = store / "evidence.json"
    if (profile_path.exists() or evidence_path.exists()) and not args.force:
        raise AgentError(
            f"事实库已存在：{store}\n"
            "如需重建，请先备份并显式传入 --force。"
        )

    source_dir = store / "source"
    source_dir.mkdir(parents=True, exist_ok=True)
    copied_resume = unique_destination(source_dir, f"resume{resume.suffix.lower()}")
    shutil.copy2(resume, copied_resume)
    resume_text = extract_text(copied_resume).strip()
    if not resume_text:
        raise AgentError("简历未提取到可读文本；扫描版 PDF 请先做 OCR。")
    resume_text_path = source_dir / "resume.txt"
    resume_text_path.write_text(resume_text + "\n", encoding="utf-8")

    profile = {
        "schema_version": SCHEMA_VERSION,
        "status": "draft",
        "candidate": {
            "name": args.candidate_name or "",
            "phone": "",
            "email": "",
            "location": "",
            "links": [],
        },
        "target_language": args.language,
        "headline": "",
        "skills": [],
        "experience": [],
        "projects": [],
        "education": [],
        "certifications": [],
    }
    evidence = {"schema_version": SCHEMA_VERSION, "items": []}
    index = {
        "schema_version": SCHEMA_VERSION,
        "created_at": now_iso(),
        "updated_at": now_iso(),
        "source_resume": str(copied_resume),
        "source_text": str(resume_text_path),
    }

    write_json(profile_path, profile)
    write_json(evidence_path, evidence)
    write_json(store / "index.json", index)
    return {
        "store": str(store),
        "profile": str(profile_path),
        "evidence": str(evidence_path),
        "source_text": str(resume_text_path),
        "next": "根据 source/resume.txt 填充 profile.json 与 evidence.json，然后完成事实确认。",
    }


def prepare_run(args: argparse.Namespace) -> dict[str, Any]:
    store = Path(args.store).expanduser().resolve()
    ensure_store(store)
    template_path = Path(args.template).expanduser().resolve()

    template_suffix = template_path.suffix.lower()
    if template_suffix not in HTML_SUFFIXES | DOCX_SUFFIXES:
        raise AgentError("模板必须是 HTML 或 DOCX；PDF 不能作为可编辑模板。")

    jd_meta: dict[str, Any] = {}
    jd_source_type = "url" if is_http_url(args.jd) else "file"
    if jd_source_type == "url":
        jd_meta = read_job_url(
            args.jd,
            store,
            headed=args.headed,
            wait_ms=args.wait_ms,
        )
        jd_text = str(jd_meta.get("text", "")).strip()
        jd_source = args.jd
    else:
        jd_path = Path(args.jd).expanduser().resolve()
        jd_suffix = jd_path.suffix.lower()
        if jd_suffix not in EXTRACTABLE_SUFFIXES:
            raise AgentError(f"不支持 JD 文件类型：{jd_suffix}")
        jd_text = extract_text(jd_path).strip()
        jd_source = str(jd_path)

    if not jd_text:
        raise AgentError("JD 未提取到可读文本。")

    company = args.company.strip() or str(jd_meta.get("inferred_company", "")).strip() or "待识别"
    role = args.role.strip() or str(jd_meta.get("inferred_role", "")).strip() or "待识别"
    filename = safe_filename_stem(args.filename or f"{company}-{role}-resume")
    timestamp = datetime.now().astimezone().strftime("%Y%m%d-%H%M%S")
    run_id = f"{timestamp}-{slugify(company)}-{slugify(role)}-{secrets.token_hex(2)}"
    run_dir = run_dir_base(store) / run_id
    inputs_dir = run_dir / "inputs"
    analysis_dir = run_dir / "analysis"
    output_dir = run_dir / "output"
    for directory in (inputs_dir, analysis_dir, output_dir):
        directory.mkdir(parents=True, exist_ok=False)

    copied_template = unique_destination(inputs_dir, f"template{template_suffix}")
    shutil.copy2(template_path, copied_template)
    if jd_source_type == "url":
        copied_jd = inputs_dir / "jd-source.json"
        write_json(copied_jd, jd_meta)
    else:
        jd_path = Path(args.jd).expanduser().resolve()
        copied_jd = unique_destination(inputs_dir, f"jd{jd_path.suffix.lower()}")
        shutil.copy2(jd_path, copied_jd)
    (inputs_dir / "jd.txt").write_text(jd_text + "\n", encoding="utf-8")
    (inputs_dir / "template.txt").write_text(extract_text(copied_template).strip() + "\n", encoding="utf-8")

    shutil.copy2(store / "profile.json", analysis_dir / "profile.json")
    shutil.copy2(store / "evidence.json", analysis_dir / "evidence.json")

    template_format = "html" if template_suffix in HTML_SUFFIXES else "docx"
    job = {
        "schema_version": SCHEMA_VERSION,
        "run_id": run_id,
        "created_at": now_iso(),
        "status": "analysis",
        "store": str(store),
        "company": company,
        "role": role,
        "language": args.language,
        "target_filename": filename,
        "input": {
            "jd": {
                "path": str(copied_jd.relative_to(run_dir)),
                "text_path": "inputs/jd.txt",
                "source": jd_source,
                "source_type": jd_source_type,
                "source_metadata_path": "inputs/jd-source.json" if jd_source_type == "url" else "",
            },
            "template": {
                "path": str(copied_template.relative_to(run_dir)),
                "text_path": "inputs/template.txt",
                "format": template_format,
                "source": str(template_path),
            },
            "profile_snapshot": "analysis/profile.json",
            "evidence_snapshot": "analysis/evidence.json",
        },
        "confirmations": {
            "facts_confirmed": False,
            "rewrites_confirmed": False,
        },
        "output": {
            "directory": "output",
            "html": f"output/{filename}.html",
            "pdf": f"output/{filename}.pdf",
            "docx": f"output/{filename}.docx" if template_format == "docx" else "",
            "change_log": "change-log.md",
            "delivery": "delivery.json",
        },
    }
    requirements = {
        "schema_version": SCHEMA_VERSION,
        "role_lens": "有多年招聘经验的 HR + 该岗位面试官",
        "company": company,
        "role": role,
        "top_capabilities": [],
        "keywords": [],
        "hr_risks": [],
        "thirty_second_review": {
            "would_interview": None,
            "decision": "",
            "where_stuck": [],
            "improvements": [],
        },
        "requirements": [],
    }
    rewrites = {"schema_version": SCHEMA_VERSION, "rewrites": []}
    write_json(run_dir / "job.json", job)
    write_json(run_dir / "requirements.json", requirements)
    write_json(run_dir / "rewrites.json", rewrites)
    return {
        "run": str(run_dir),
        "job": str(run_dir / "job.json"),
        "requirements": str(run_dir / "requirements.json"),
        "rewrites": str(run_dir / "rewrites.json"),
        "company": company,
        "role": role,
        "jd_source_type": jd_source_type,
        "next": (
            "确认公司和岗位名称，分析 JD 和事实，填写 requirements.json 与 "
            "rewrites.json，再进入两个确认 Gate。"
        ),
    }


def promote_facts(args: argparse.Namespace) -> dict[str, Any]:
    run_dir, job, _profile, _evidence, _requirements, _rewrites, _result = validate_run(
        Path(args.run),
        strict_confirmations=False,
    )
    if (job.get("confirmations") or {}).get("facts_confirmed") is not True:
        raise AgentError("事实确认 Gate 尚未完成，不能同步到长期事实库。")

    store_value = job.get("store")
    if not store_value:
        raise AgentError("job.json 缺少 store 路径，无法同步事实库。")
    store = Path(store_value).expanduser().resolve()
    store.mkdir(parents=True, exist_ok=True)
    source_profile = resolve_run_file(run_dir, job["input"]["profile_snapshot"])
    source_evidence = resolve_run_file(run_dir, job["input"]["evidence_snapshot"])
    shutil.copy2(source_profile, store / "profile.json")
    shutil.copy2(source_evidence, store / "evidence.json")

    index_path = store / "index.json"
    if index_path.exists():
        index = read_json(index_path)
    else:
        index = {"schema_version": SCHEMA_VERSION, "created_at": now_iso()}
    index["updated_at"] = now_iso()
    index["last_promoted_run"] = job.get("run_id", "")
    write_json(index_path, index)
    return {
        "store": str(store),
        "profile": str(store / "profile.json"),
        "evidence": str(store / "evidence.json"),
        "run": str(run_dir),
        "next": "继续填写 requirements.json 与 rewrites.json，并完成改写确认 Gate。",
    }


def resolve_run_file(run_dir: Path, value: str) -> Path:
    path = Path(value)
    return path if path.is_absolute() else run_dir / path


def load_run(run_dir: Path) -> tuple[Path, dict[str, Any], dict[str, Any], dict[str, Any], dict[str, Any], dict[str, Any]]:
    run_dir = run_dir.expanduser().resolve()
    job = read_json(run_dir / "job.json")
    profile = read_json(resolve_run_file(run_dir, job["input"]["profile_snapshot"]))
    evidence = read_json(resolve_run_file(run_dir, job["input"]["evidence_snapshot"]))
    requirements = read_json(run_dir / "requirements.json")
    rewrites = read_json(run_dir / "rewrites.json")
    return run_dir, job, profile, evidence, requirements, rewrites


def validate_run(
    run_dir: Path,
    strict_confirmations: bool = False,
) -> tuple[
    Path,
    dict[str, Any],
    dict[str, Any],
    dict[str, Any],
    dict[str, Any],
    dict[str, Any],
    dict[str, Any],
]:
    run_dir, job, profile, evidence, requirements, rewrites = load_run(run_dir)
    errors: list[str] = []
    warnings: list[str] = []

    if job.get("schema_version") != SCHEMA_VERSION:
        errors.append("job.json 的 schema_version 不受支持。")
    company = str(job.get("company", "")).strip()
    role = str(job.get("role", "")).strip()
    if not company or not role:
        errors.append("job.json 必须包含 company 和 role。")
    elif company in UNRESOLVED_METADATA or role in UNRESOLVED_METADATA:
        if strict_confirmations:
            errors.append("最终交付前必须确认公司和岗位名称，不能保留“待识别”。")
        else:
            warnings.append("公司和岗位名称尚未完成确认。")
    jd_info = job.get("input", {}).get("jd", {})
    if jd_info.get("source_type") not in {None, "file", "url"}:
        errors.append("job.json 的 JD source_type 必须是 file 或 url。")
    template_info = job.get("input", {}).get("template", {})
    template_format = template_info.get("format")
    if template_format not in {"html", "docx"}:
        errors.append("job.json 的模板 format 必须是 html 或 docx。")
    else:
        template_path = resolve_run_file(run_dir, template_info.get("path", ""))
        if not template_path.exists():
            errors.append(f"模板不存在：{template_path}")
        if template_format == "html" and template_path.suffix.lower() not in HTML_SUFFIXES:
            errors.append("HTML 模板扩展名无效。")
        if template_format == "docx" and template_path.suffix.lower() not in DOCX_SUFFIXES:
            errors.append("DOCX 模板扩展名无效。")

    if profile.get("schema_version") != SCHEMA_VERSION:
        errors.append("profile.json 的 schema_version 不受支持。")
    candidate = profile.get("candidate", {})
    if not isinstance(candidate, dict) or not candidate.get("name"):
        errors.append("profile.candidate.name 不能为空。")

    evidence_items = evidence.get("items")
    if not isinstance(evidence_items, list):
        errors.append("evidence.json.items 必须是数组。")
        evidence_items = []
    evidence_by_id: dict[str, dict[str, Any]] = {}
    allowed_statuses = {"verified", "unverified", "conflicting"}
    allowed_ownership = {"led", "owned", "participated", "supported", "observed"}
    for item in evidence_items:
        evidence_id = item.get("id") if isinstance(item, dict) else None
        if not evidence_id:
            errors.append("存在缺少 id 的 EvidenceItem。")
            continue
        if evidence_id in evidence_by_id:
            errors.append(f"EvidenceItem id 重复：{evidence_id}")
        evidence_by_id[evidence_id] = item
        if item.get("status") not in allowed_statuses:
            errors.append(f"{evidence_id}: status 必须是 verified/unverified/conflicting。")
        if item.get("ownership") not in allowed_ownership:
            errors.append(f"{evidence_id}: ownership 无效。")
        if not str(item.get("statement", "")).strip():
            errors.append(f"{evidence_id}: statement 不能为空。")
        if not item.get("source_refs"):
            errors.append(f"{evidence_id}: 至少需要一个 source_refs。")
        if item.get("status") == "verified" and item.get("ownership") in {"participated", "supported", "observed"}:
            prohibited = item.get("prohibited_phrases") or []
            if not prohibited:
                warnings.append(f"{evidence_id}: 非主导经历建议填写 prohibited_phrases。")

    requirement_items = requirements.get("requirements")
    if not isinstance(requirement_items, list):
        errors.append("requirements.json.requirements 必须是数组。")
        requirement_items = []
    requirement_ids: set[str] = set()
    requirements_by_id: dict[str, dict[str, Any]] = {}
    for requirement in requirement_items:
        requirement_id = requirement.get("id") if isinstance(requirement, dict) else None
        if not requirement_id:
            errors.append("存在缺少 id 的 JDRequirement。")
            continue
        if requirement_id in requirement_ids:
            errors.append(f"JDRequirement id 重复：{requirement_id}")
        requirement_ids.add(requirement_id)
        requirements_by_id[requirement_id] = requirement
        if requirement.get("priority") not in {"must", "preferred"}:
            errors.append(f"{requirement_id}: priority 必须是 must 或 preferred。")
        if requirement.get("coverage") not in {"strong", "transferable", "gap"}:
            errors.append(f"{requirement_id}: coverage 无效。")
        if not str(requirement.get("requirement", "")).strip():
            errors.append(f"{requirement_id}: requirement 不能为空。")
        linked = requirement.get("evidence_ids") or []
        if requirement.get("coverage") != "gap" and not linked:
            errors.append(f"{requirement_id}: 非 gap 要求必须关联证据。")
        if requirement.get("coverage") == "gap" and not str(requirement.get("gap_note", "")).strip():
            errors.append(f"{requirement_id}: gap 要求必须填写 gap_note。")
        for evidence_id in linked:
            if evidence_id not in evidence_by_id:
                errors.append(f"{requirement_id}: 引用了不存在的证据 {evidence_id}。")

    top_capabilities = requirements.get("top_capabilities")
    if not isinstance(top_capabilities, list):
        errors.append("requirements.json.top_capabilities 必须是数组。")
        top_capabilities = []
    elif any(not isinstance(item, str) or not item.strip() for item in top_capabilities):
        errors.append("requirements.json.top_capabilities 只能包含非空字符串。")
    keywords = requirements.get("keywords")
    if not isinstance(keywords, list):
        errors.append("requirements.json.keywords 必须是数组。")
        keywords = []
    elif any(not isinstance(item, str) or not item.strip() for item in keywords):
        errors.append("requirements.json.keywords 只能包含非空字符串。")
    hr_risks = requirements.get("hr_risks")
    if not isinstance(hr_risks, list):
        errors.append("requirements.json.hr_risks 必须是数组。")
        hr_risks = []
    elif any(not isinstance(item, str) or not item.strip() for item in hr_risks):
        errors.append("requirements.json.hr_risks 只能包含非空字符串。")
    review = requirements.get("thirty_second_review")
    if not isinstance(review, dict):
        errors.append("requirements.json.thirty_second_review 必须是对象。")
        review = {}
    else:
        would_interview = review.get("would_interview")
        if would_interview is not None and not isinstance(would_interview, bool):
            errors.append("thirty_second_review.would_interview 必须是布尔值或 null。")
        if not str(review.get("decision", "")).strip():
            warnings.append("thirty_second_review.decision 尚未填写。")
        for field in ("where_stuck", "improvements"):
            value = review.get(field)
            if not isinstance(value, list):
                errors.append(f"thirty_second_review.{field} 必须是数组。")
            elif any(not isinstance(item, str) or not item.strip() for item in value):
                errors.append(f"thirty_second_review.{field} 只能包含非空字符串。")
        if would_interview is False:
            if not review.get("where_stuck"):
                errors.append("30 秒复核未通过时，必须填写 where_stuck。")
            if not review.get("improvements"):
                errors.append("30 秒复核未通过时，必须填写 improvements。")

    rewrite_items = rewrites.get("rewrites")
    if not isinstance(rewrite_items, list):
        errors.append("rewrites.json.rewrites 必须是数组。")
        rewrite_items = []
    rewrite_ids: set[str] = set()
    approved_targets: set[str] = set()
    proposed_count = 0
    approved_count = 0
    rejected_count = 0
    for rewrite in rewrite_items:
        rewrite_id = rewrite.get("id") if isinstance(rewrite, dict) else None
        if not rewrite_id:
            errors.append("存在缺少 id 的 ResumeRewrite。")
            continue
        if rewrite_id in rewrite_ids:
            errors.append(f"ResumeRewrite id 重复：{rewrite_id}")
        rewrite_ids.add(rewrite_id)
        status = rewrite.get("status")
        if status not in {"proposed", "approved", "rejected"}:
            errors.append(f"{rewrite_id}: status 无效。")
            continue
        if status == "proposed":
            proposed_count += 1
        elif status == "approved":
            approved_count += 1
        else:
            rejected_count += 1

        for field in ("original", "proposed", "rationale"):
            if not str(rewrite.get(field, "")).strip():
                errors.append(f"{rewrite_id}: {field} 不能为空。")

        target = rewrite.get("target")
        if not isinstance(target, dict):
            errors.append(f"{rewrite_id}: target 必须是对象。")
            target_key = ""
        elif template_format == "html":
            xpath = str(target.get("xpath", "")).strip()
            if not xpath:
                errors.append(f"{rewrite_id}: HTML target 缺少 xpath。")
            target_key = xpath
        elif template_format == "docx":
            if not isinstance(target.get("paragraph_index"), int):
                errors.append(f"{rewrite_id}: DOCX target 缺少 paragraph_index。")
            target_key = json.dumps(
                {key: target.get(key) for key in ("table_index", "row_index", "cell_index", "paragraph_index", "run_index")},
                sort_keys=True,
                ensure_ascii=False,
            )
        else:
            target_key = ""

        linked_requirement_ids = rewrite.get("requirement_ids") or []
        if not isinstance(linked_requirement_ids, list):
            errors.append(f"{rewrite_id}: requirement_ids 必须是数组。")
            linked_requirement_ids = []
        for requirement_id in linked_requirement_ids:
            if requirement_id not in requirements_by_id:
                errors.append(f"{rewrite_id}: 引用了不存在的 JD 要求 {requirement_id}。")

        linked_ids = rewrite.get("evidence_ids") or []
        if not isinstance(linked_ids, list):
            errors.append(f"{rewrite_id}: evidence_ids 必须是数组。")
            linked_ids = []
        if not linked_ids:
            if status == "approved":
                errors.append(f"{rewrite_id}: 获批改写必须关联证据。")
            continue

        linked_evidence: list[dict[str, Any]] = []
        for evidence_id in linked_ids:
            item = evidence_by_id.get(evidence_id)
            if item is None:
                errors.append(f"{rewrite_id}: 引用了不存在的证据 {evidence_id}。")
            else:
                linked_evidence.append(item)

        if status != "approved":
            continue

        if target_key:
            if target_key in approved_targets:
                errors.append(f"{rewrite_id}: 同一模板位置存在多个获批改写。")
            approved_targets.add(target_key)

        if not linked_evidence:
            continue
        if any(item.get("status") != "verified" for item in linked_evidence):
            errors.append(f"{rewrite_id}: 获批改写只能引用 verified 证据。")
            continue

        proposed = str(rewrite.get("proposed", ""))
        if PLACEHOLDER_RE.search(proposed):
            errors.append(f"{rewrite_id}: 获批文本仍包含占位符。")
        if VAGUE_CLAIM_RE.search(proposed):
            errors.append(f"{rewrite_id}: 空泛形容词未转成可验证证据。")

        valid_requirement_chain = False
        for requirement_id in linked_requirement_ids:
            requirement = requirements_by_id.get(requirement_id)
            if not requirement or requirement.get("coverage") == "gap":
                continue
            requirement_evidence = set(requirement.get("evidence_ids") or [])
            if requirement_evidence.intersection(linked_ids):
                valid_requirement_chain = True
                break
        if not linked_requirement_ids:
            errors.append(f"{rewrite_id}: 获批改写必须关联至少一条 JD 要求。")
        elif not valid_requirement_chain:
            errors.append(f"{rewrite_id}: 未形成“JD 要求 -> 证据 -> 改写”的完整链路。")

        evidence_blob = json.dumps(linked_evidence, ensure_ascii=False)
        original = str(rewrite.get("original", ""))
        for number in NUMBER_RE.findall(proposed):
            if number in original:
                continue
            if number not in evidence_blob:
                errors.append(f"{rewrite_id}: 新数字 {number} 无法在关联证据中定位。")

        if STRONG_CLAIM_RE.search(proposed):
            weak_ownership = any(
                item.get("ownership") in {"participated", "supported", "observed"}
                for item in linked_evidence
            )
            if weak_ownership:
                errors.append(f"{rewrite_id}: 改写扩大了角色边界，出现强责任用语。")

        for item in linked_evidence:
            for phrase in item.get("prohibited_phrases") or []:
                if phrase and phrase in proposed:
                    errors.append(f"{rewrite_id}: 命中证据禁止短语“{phrase}”。")

        if not any(item.get("metrics") for item in linked_evidence) and not RESULT_SIGNAL_RE.search(proposed):
            warnings.append(
                f"{rewrite_id}: 未体现数字、规模、频率、范围、排名、对比或结果信号，建议补强证据强度。"
            )

    if strict_confirmations:
        confirmations = job.get("confirmations") or {}
        if confirmations.get("facts_confirmed") is not True:
            errors.append("尚未完成事实确认 Gate：job.confirmations.facts_confirmed 必须为 true。")
        if confirmations.get("rewrites_confirmed") is not True:
            errors.append("尚未完成改写确认 Gate：job.confirmations.rewrites_confirmed 必须为 true。")
        if proposed_count:
            errors.append(f"仍有 {proposed_count} 条 proposed 改写，未完成 Gate 2。")
        if not 3 <= len(top_capabilities) <= 5:
            errors.append("最终交付前必须提炼 3-5 项岗位最看重的能力。")
        if len(keywords) < 3:
            errors.append("最终交付前必须提炼至少 3 个 JD 关键词。")
        if not hr_risks:
            errors.append("最终交付前必须记录至少一个 HR 可能跳过的风险或判断。")
        if review.get("would_interview") is not True and review.get("would_interview") is not False:
            errors.append("最终交付前必须完成 30 秒 HR 红队复核。")
        if not str(review.get("decision", "")).strip():
            errors.append("最终交付前必须填写 30 秒 HR 复核结论。")

    if not requirement_items:
        warnings.append("requirements.json 为空，JD 匹配分析尚未完成。")
    if not rewrite_items:
        warnings.append("rewrites.json 为空，尚无岗位定制改写。")

    result = {
        "ok": not errors,
        "errors": errors,
        "warnings": warnings,
        "counts": {
            "evidence": len(evidence_items),
            "verified_evidence": sum(1 for item in evidence_items if item.get("status") == "verified"),
            "requirements": len(requirement_items),
            "rewrites": len(rewrite_items),
            "approved": approved_count,
            "rejected": rejected_count,
            "proposed": proposed_count,
        },
    }
    return run_dir, job, profile, evidence, requirements, rewrites, result


def print_validation(run_dir: Path, strict_confirmations: bool) -> bool:
    *_, result = validate_run(run_dir, strict_confirmations)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return bool(result["ok"])


def resolve_docx_paragraph(document: Any, target: dict[str, Any]) -> Any:
    if "table_index" in target:
        table_index = int(target["table_index"])
        row_index = int(target.get("row_index", 0))
        cell_index = int(target.get("cell_index", 0))
        paragraph_index = int(target.get("paragraph_index", 0))
        try:
            cell = document.tables[table_index].rows[row_index].cells[cell_index]
            return cell.paragraphs[paragraph_index]
        except (IndexError, TypeError) as exc:
            raise AgentError(f"DOCX 表格定位越界：{target}") from exc
    paragraph_index = int(target.get("paragraph_index", -1))
    if paragraph_index < 0 or paragraph_index >= len(document.paragraphs):
        raise AgentError(f"DOCX 段落定位越界：{paragraph_index}")
    return document.paragraphs[paragraph_index]


def set_docx_text(paragraph: Any, proposed: str, run_index: Any = None) -> None:
    if run_index is not None:
        index = int(run_index)
        if index < 0 or index >= len(paragraph.runs):
            raise AgentError(f"DOCX run_index 越界：{index}")
        paragraph.runs[index].text = proposed
        return
    if not paragraph.runs:
        paragraph.add_run(proposed)
        return
    paragraph.runs[0].text = proposed
    for run in paragraph.runs[1:]:
        run.text = ""


def replace_html_node(node: Any, proposed: str, mode: str) -> None:
    if mode not in {"text", "html"}:
        raise AgentError(f"不支持的 HTML target.mode：{mode}")
    if mode == "text":
        node.text = proposed
        for child in list(node):
            node.remove(child)
        return

    from lxml import html

    fragment = html.fragment_fromstring(f"<span>{proposed}</span>")
    node.text = fragment.text
    for child in list(node):
        node.remove(child)
    for child in list(fragment):
        node.append(child)


def apply_rewrites(
    run_dir: Path,
    job: dict[str, Any],
    rewrites: dict[str, Any],
) -> tuple[Path, str]:
    template_info = job["input"]["template"]
    template_format = template_info["format"]
    source = resolve_run_file(run_dir, template_info["path"])
    output_dir = run_dir / job["output"]["directory"]
    output_dir.mkdir(parents=True, exist_ok=True)
    target_stem = safe_filename_stem(job["target_filename"])

    approved = [
        rewrite for rewrite in rewrites.get("rewrites", [])
        if rewrite.get("status") == "approved"
    ]

    if template_format == "html":
        from lxml import etree, html

        output_path = output_dir / f"{target_stem}.html"
        document = html.parse(str(source))
        seen_xpaths: set[str] = set()
        for rewrite in approved:
            target = rewrite.get("target") or {}
            xpath = str(target.get("xpath", "")).strip()
            if not xpath.startswith("/"):
                xpath = "//" + xpath.lstrip("/")
            if xpath in seen_xpaths:
                raise AgentError(f"{rewrite['id']}: 同一 XPath 被重复改写。")
            seen_xpaths.add(xpath)
            matches = document.xpath(xpath)
            if len(matches) != 1:
                raise AgentError(f"{rewrite['id']}: XPath 必须唯一匹配，实际匹配 {len(matches)} 个。")
            node = matches[0]
            if not hasattr(node, "tag"):
                raise AgentError(f"{rewrite['id']}: XPath 必须指向元素，不能只指向文本。")
            ancestor_classes = " ".join(
                str(ancestor.get("class", ""))
                for ancestor in [node, *node.iterancestors()]
                if hasattr(ancestor, "get")
            )
            if "no-print" in ancestor_classes.split():
                raise AgentError(f"{rewrite['id']}: 不允许改写 no-print 区域。")
            current = " ".join(node.text_content().split())
            original = " ".join(str(rewrite.get("original", "")).split())
            if current and original and current != original:
                raise AgentError(
                    f"{rewrite['id']}: 原文与模板当前文本不一致。\n"
                    f"模板：{current}\nJSON：{original}"
                )
            replace_html_node(node, str(rewrite["proposed"]), str(target.get("mode", "text")))

        serialized = etree.tostring(
            document,
            method="html",
            encoding="unicode",
            doctype="<!DOCTYPE html>",
        )
        output_path.write_text(serialized + "\n", encoding="utf-8")
        return output_path, "html"

    if template_format == "docx":
        try:
            from docx import Document
            from docx.enum.section import WD_ORIENT
            from docx.shared import Mm
        except ImportError as exc:
            raise AgentError("改写 DOCX 需要 python-docx。") from exc
        output_path = output_dir / f"{target_stem}.docx"
        shutil.copy2(source, output_path)
        document = Document(str(output_path))
        seen_targets: set[str] = set()
        for rewrite in approved:
            target = rewrite.get("target") or {}
            target_key = json.dumps(target, sort_keys=True, ensure_ascii=False)
            if target_key in seen_targets:
                raise AgentError(f"{rewrite['id']}: 同一 DOCX 位置被重复改写。")
            seen_targets.add(target_key)
            paragraph = resolve_docx_paragraph(document, target)
            current = " ".join(paragraph.text.split())
            original = " ".join(str(rewrite.get("original", "")).split())
            if current and original and current != original:
                raise AgentError(
                    f"{rewrite['id']}: 原文与 DOCX 当前文本不一致。\n"
                    f"模板：{current}\nJSON：{original}"
                )
            set_docx_text(paragraph, str(rewrite["proposed"]), target.get("run_index"))
        for section in document.sections:
            if section.orientation == WD_ORIENT.LANDSCAPE:
                section.page_width = Mm(297)
                section.page_height = Mm(210)
            else:
                section.page_width = Mm(210)
                section.page_height = Mm(297)
        document.save(str(output_path))
        return output_path, "docx"

    raise AgentError(f"不支持的模板格式：{template_format}")


def find_node_executable() -> str:
    candidates = [
        os.environ.get("CODEX_NODE"),
        shutil.which("node"),
    ]
    home = Path.home()
    candidates.extend(
        str(path)
        for path in sorted(
            home.glob(".cache/codex-runtimes/*/dependencies/node/bin/node"),
            reverse=True,
        )
    )
    for candidate in candidates:
        if candidate and Path(candidate).is_file():
            return str(candidate)
    raise AgentError("找不到 Node.js，无法导出 HTML PDF。")


def find_node_modules() -> Path | None:
    env_path = os.environ.get("NODE_PATH")
    if env_path:
        for entry in env_path.split(os.pathsep):
            candidate = Path(entry)
            if (candidate / "playwright").exists():
                return candidate
    home = Path.home()
    for candidate in sorted(
        home.glob(".cache/codex-runtimes/*/dependencies/node/node_modules"),
        reverse=True,
    ):
        if (candidate / "playwright").exists():
            return candidate
    return None


def find_chrome_executable() -> str | None:
    candidates = [
        os.environ.get("CHROME_PATH"),
        "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
        "/Applications/Chromium.app/Contents/MacOS/Chromium",
        "/Applications/Microsoft Edge.app/Contents/MacOS/Microsoft Edge",
        shutil.which("google-chrome"),
        shutil.which("chromium"),
        shutil.which("chromium-browser"),
    ]
    for candidate in candidates:
        if candidate and Path(candidate).is_file():
            return str(candidate)
    return None


def find_soffice_executable() -> str:
    candidates = [
        os.environ.get("SOFFICE_PATH"),
        shutil.which("soffice"),
    ]
    home = Path.home()
    candidates.extend(
        str(path)
        for path in sorted(
            home.glob(".cache/codex-runtimes/*/dependencies/bin/*/soffice"),
            reverse=True,
        )
    )
    for candidate in candidates:
        if candidate and Path(candidate).is_file():
            return str(candidate)
    raise AgentError("找不到 LibreOffice/soffice，无法从 DOCX 导出 PDF。")


def read_job_url(
    url: str,
    store: Path,
    headed: bool = False,
    wait_ms: int = 0,
) -> dict[str, Any]:
    script = Path(__file__).with_name("read_job_page.mjs")
    if not script.exists():
        raise AgentError(f"缺少岗位网页读取脚本：{script}")
    node = find_node_executable()
    env = os.environ.copy()
    node_modules = find_node_modules()
    if node_modules:
        env["NODE_PATH"] = str(node_modules)
    chrome = find_chrome_executable()
    if chrome:
        env["CHROME_PATH"] = chrome

    effective_wait = wait_ms if wait_ms > 0 else (30_000 if headed else 5_000)
    command = [
        node,
        str(script),
        "--url",
        url,
        "--profile-dir",
        str(store / "browser-profile"),
        "--wait-ms",
        str(effective_wait),
    ]
    if headed:
        command.append("--headed")
    process = subprocess.run(
        command,
        capture_output=True,
        text=True,
        env=env,
    )

    payload: dict[str, Any] | None = None
    for line in reversed(process.stdout.splitlines()):
        try:
            parsed = json.loads(line)
        except json.JSONDecodeError:
            continue
        if isinstance(parsed, dict):
            payload = parsed
            break
    if payload is None:
        detail = (process.stderr or process.stdout or "").strip()
        raise AgentError(f"岗位链接读取失败：\n{detail}")
    if process.returncode != 0 or payload.get("ok") is not True:
        detail = payload.get("error") or (process.stderr or process.stdout or "").strip()
        if payload.get("needs_login"):
            detail += (
                "\n请使用 `--headed --wait-ms 60000` 重新运行，在打开的浏览器中登录或完成验证后等待自动提取。"
            )
        raise AgentError(f"岗位链接读取失败：{detail}")
    return payload


def export_html_pdf(html_path: Path, pdf_path: Path) -> dict[str, Any]:
    script = Path(__file__).with_name("export_pdf.mjs")
    if not script.exists():
        raise AgentError(f"缺少 PDF 导出脚本：{script}")
    node = find_node_executable()
    env = os.environ.copy()
    node_modules = find_node_modules()
    if node_modules:
        env["NODE_PATH"] = str(node_modules)
    chrome = find_chrome_executable()
    if chrome:
        env["CHROME_PATH"] = chrome
    process = subprocess.run(
        [node, str(script), str(html_path), str(pdf_path)],
        capture_output=True,
        text=True,
        env=env,
    )
    if process.returncode != 0:
        detail = (process.stderr or process.stdout or "").strip()
        raise AgentError(f"HTML PDF 导出失败：\n{detail}")
    metrics: dict[str, Any] | None = None
    for line in reversed(process.stdout.splitlines()):
        try:
            parsed = json.loads(line)
        except json.JSONDecodeError:
            continue
        if isinstance(parsed, dict):
            metrics = parsed
            break
    if metrics is None:
        raise AgentError(f"PDF 导出脚本没有返回渲染指标：\n{process.stdout}")
    return metrics


def convert_docx(docx_path: Path, target_format: str, output_dir: Path) -> Path:
    soffice = find_soffice_executable()
    with tempfile.TemporaryDirectory(prefix="resume-agent-libreoffice-") as profile_dir:
        profile_uri = Path(profile_dir).resolve().as_uri()
        command = [
            soffice,
            f"-env:UserInstallation={profile_uri}",
            "--headless",
            "--convert-to",
            target_format,
            "--outdir",
            str(output_dir),
            str(docx_path),
        ]
        process = subprocess.run(command, capture_output=True, text=True)
    if process.returncode != 0:
        detail = (process.stderr or process.stdout or "").strip()
        raise AgentError(f"DOCX 转换 {target_format} 失败：\n{detail}")
    suffix = ".html" if target_format == "html" else ".pdf"
    converted = output_dir / f"{docx_path.stem}{suffix}"
    if not converted.exists():
        raise AgentError(f"DOCX 转换后未找到文件：{converted}")
    return converted


def inspect_pdf(pdf_path: Path) -> dict[str, Any]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise AgentError("检查 PDF 需要 pypdf，请使用 Codex workspace dependency 的 Python。") from exc

    reader = PdfReader(str(pdf_path))
    page_count = len(reader.pages)
    page_sizes: list[list[float]] = []
    a4_ok = True
    for page in reader.pages:
        box = page.mediabox
        width = float(box.width)
        height = float(box.height)
        page_sizes.append([round(width, 2), round(height, 2)])
        short = min(width, height)
        long = max(width, height)
        if abs(short - 595.28) > 3 or abs(long - 841.89) > 3:
            a4_ok = False

    text = ""
    try:
        import pdfplumber

        logging.getLogger("pdfminer").setLevel(logging.ERROR)
        with pdfplumber.open(str(pdf_path)) as pdf:
            text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    except ImportError:
        text = ""

    return {
        "page_count": page_count,
        "page_sizes_points": page_sizes,
        "a4": a4_ok,
        "text_length": len(text.strip()),
        "text_selectable": len(text.strip()) >= 20,
    }


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def markdown_cell(value: Any) -> str:
    return str(value or "").replace("|", "\\|").replace("\n", "<br>")


def build_change_log(
    job: dict[str, Any],
    evidence: dict[str, Any],
    requirements: dict[str, Any],
    rewrites: dict[str, Any],
    validation: dict[str, Any],
    pdf_metrics: dict[str, Any],
) -> str:
    evidence_by_id = {
        item["id"]: item for item in evidence.get("items", []) if item.get("id")
    }
    lines = [
        "# 简历定制修改记录",
        "",
        "> 简历价值 = 匹配度 × 证据强度 × 可读性 ÷ 对方决策成本。",
        "> 招聘方视角：有多年招聘经验的 HR + 该岗位面试官。",
        "",
        f"- 公司：{job.get('company', '')}",
        f"- 岗位：{job.get('role', '')}",
        f"- 生成时间：{now_iso()}",
        f"- JD：{job.get('input', {}).get('jd', {}).get('source', '')}",
        f"- 模板：{job.get('input', {}).get('template', {}).get('path', '')}",
        f"- 事实确认：{job.get('confirmations', {}).get('facts_confirmed')}",
        f"- 改写确认：{job.get('confirmations', {}).get('rewrites_confirmed')}",
        f"- PDF 页数：{pdf_metrics.get('page_count')}",
        f"- A4：{pdf_metrics.get('a4')}",
        f"- 文字可提取：{pdf_metrics.get('text_selectable')}",
        "",
        "## 岗位判断",
        "",
        f"- 最看重能力：{'、'.join(requirements.get('top_capabilities') or [])}",
        f"- 核心关键词：{'、'.join(requirements.get('keywords') or [])}",
        f"- HR 风险：{'；'.join(requirements.get('hr_risks') or [])}",
        f"- 30 秒复核：{requirements.get('thirty_second_review', {}).get('decision', '')}",
        "",
        "## 逐条修改",
        "",
        "| ID | 状态 | JD 要求 | 原文 | 新描述 | 证据 | 修改理由 |",
        "| --- | --- | --- | --- | --- | --- | --- |",
    ]
    for rewrite in rewrites.get("rewrites", []):
        evidence_labels = []
        for evidence_id in rewrite.get("evidence_ids") or []:
            item = evidence_by_id.get(evidence_id, {})
            ownership = item.get("ownership", "")
            evidence_labels.append(f"{evidence_id} ({ownership})" if ownership else evidence_id)
        requirement_ids = ", ".join(rewrite.get("requirement_ids") or [])
        lines.append(
            "| {id} | {status} | {requirements} | {original} | {proposed} | {evidence} | {rationale} |".format(
                id=markdown_cell(rewrite.get("id")),
                status=markdown_cell(rewrite.get("status")),
                requirements=markdown_cell(requirement_ids),
                original=markdown_cell(rewrite.get("original")),
                proposed=markdown_cell(rewrite.get("proposed")),
                evidence=markdown_cell(", ".join(evidence_labels)),
                rationale=markdown_cell(rewrite.get("rationale")),
            )
        )
    if not rewrites.get("rewrites"):
        lines.append("| - | - | - | 未生成改写 | - | - | - |")
    lines.extend(
        [
            "",
            "## 校验摘要",
            "",
            f"- 结构校验：{'通过' if validation.get('ok') else '失败'}",
            f"- 验证后证据数：{validation.get('counts', {}).get('verified_evidence', 0)}",
            f"- 获批改写数：{validation.get('counts', {}).get('approved', 0)}",
            f"- 驳回改写数：{validation.get('counts', {}).get('rejected', 0)}",
            "",
        ]
    )
    return "\n".join(lines)


def finalize_run(args: argparse.Namespace) -> dict[str, Any]:
    run_dir = Path(args.run).expanduser().resolve()
    run_dir, job, _profile, evidence, requirements, rewrites_data, validation = validate_run(
        run_dir,
        strict_confirmations=True,
    )
    if not validation["ok"]:
        errors = "\n".join(f"- {error}" for error in validation["errors"])
        raise AgentError(f"导出前校验失败：\n{errors}")

    applied_path, applied_format = apply_rewrites(run_dir, job, rewrites_data)
    output_dir = run_dir / job["output"]["directory"]
    target_stem = safe_filename_stem(job["target_filename"])

    render_metrics: dict[str, Any] = {}
    if applied_format == "html":
        html_path = applied_path
        pdf_path = output_dir / f"{target_stem}.pdf"
        render_metrics = export_html_pdf(html_path, pdf_path)
    else:
        docx_path = applied_path
        pdf_path = convert_docx(docx_path, "pdf", output_dir)
        html_candidate = convert_docx(docx_path, "html", output_dir)
        html_path = output_dir / f"{target_stem}.html"
        if html_candidate != html_path:
            if html_path.exists():
                html_path.unlink()
            html_candidate.replace(html_path)

    pdf_metrics = inspect_pdf(pdf_path)
    pdf_errors: list[str] = []
    if pdf_metrics["page_count"] != 1:
        pdf_errors.append(f"PDF 必须恰好一页，当前为 {pdf_metrics['page_count']} 页。")
    if not pdf_metrics["a4"]:
        pdf_errors.append(f"PDF 不是 A4：{pdf_metrics['page_sizes_points']}")
    if not pdf_metrics["text_selectable"]:
        pdf_errors.append("PDF 文本无法可靠提取，可能被栅格化。")
    if render_metrics.get("horizontal_overflow"):
        pdf_errors.append("HTML 存在水平溢出。")

    change_log = build_change_log(
        job,
        evidence,
        requirements,
        rewrites_data,
        validation,
        pdf_metrics,
    )
    change_log_path = run_dir / "change-log.md"
    change_log_path.write_text(change_log, encoding="utf-8")

    delivery = {
        "schema_version": SCHEMA_VERSION,
        "run_id": job["run_id"],
        "created_at": now_iso(),
        "status": "complete" if not pdf_errors else "failed",
        "inputs": job["input"],
        "outputs": {
            "html": str(html_path.relative_to(run_dir)),
            "pdf": str(pdf_path.relative_to(run_dir)),
            "docx": str(applied_path.relative_to(run_dir)) if applied_format == "docx" else "",
            "change_log": "change-log.md",
        },
        "hashes": {
            "html": sha256_file(html_path),
            "pdf": sha256_file(pdf_path),
            "change_log": sha256_file(change_log_path),
        },
        "render_metrics": render_metrics,
        "pdf_metrics": pdf_metrics,
        "validation": validation,
        "errors": pdf_errors,
    }
    if applied_format == "docx":
        delivery["hashes"]["docx"] = sha256_file(applied_path)
    write_json(run_dir / "delivery.json", delivery)
    if pdf_errors:
        raise AgentError("交付校验失败：\n" + "\n".join(f"- {error}" for error in pdf_errors))

    job["status"] = "complete"
    job["completed_at"] = now_iso()
    write_json(run_dir / "job.json", job)
    return {
        "run": str(run_dir),
        "html": str(html_path),
        "pdf": str(pdf_path),
        "docx": str(applied_path) if applied_format == "docx" else "",
        "change_log": str(change_log_path),
        "delivery": str(run_dir / "delivery.json"),
        "pdf_metrics": pdf_metrics,
        "render_metrics": render_metrics,
    }


def status_run(args: argparse.Namespace) -> dict[str, Any]:
    run_dir, job, *_rest = load_run(Path(args.run))
    return {
        "run": str(run_dir),
        "status": job.get("status"),
        "company": job.get("company"),
        "role": job.get("role"),
        "confirmations": job.get("confirmations"),
        "output": job.get("output"),
    }


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="本地简历事实库、岗位定制改写、A4 校验与 PDF 导出工具。",
    )
    subparsers = parser.add_subparsers(dest="command", required=True)

    init_parser = subparsers.add_parser("init", help="初始化本地事实库并提取简历文本。")
    init_parser.add_argument("--store", default=".resume-agent", help="事实库目录。")
    init_parser.add_argument("--resume", required=True, help="已有简历文件。")
    init_parser.add_argument("--candidate-name", default="", help="候选人姓名。")
    init_parser.add_argument("--language", default="zh-CN", help="目标简历语言。")
    init_parser.add_argument("--force", action="store_true", help="覆盖已有 profile 和 evidence。")
    init_parser.set_defaults(handler=init_store)

    prepare_parser = subparsers.add_parser("prepare", help="为一次投递创建独立 run。")
    prepare_parser.add_argument("--store", default=".resume-agent", help="事实库目录。")
    prepare_parser.add_argument("--company", default="", help="公司名称；留空时尝试从链接提取。")
    prepare_parser.add_argument("--role", default="", help="岗位名称；留空时尝试从链接提取。")
    prepare_parser.add_argument("--jd", required=True, help="JD 文件路径或投递链接。")
    prepare_parser.add_argument("--template", required=True, help="HTML 或 DOCX 简历模板。")
    prepare_parser.add_argument("--filename", default="", help="输出文件名主干。")
    prepare_parser.add_argument("--language", default="zh-CN", help="目标简历语言。")
    prepare_parser.add_argument(
        "--headed",
        action="store_true",
        help="打开可见浏览器，用于登录或完成验证码后读取岗位链接。",
    )
    prepare_parser.add_argument(
        "--wait-ms",
        type=int,
        default=0,
        help="页面加载后等待毫秒数；默认无头 5000，可见浏览器 30000。",
    )
    prepare_parser.set_defaults(handler=prepare_run)

    validate_parser = subparsers.add_parser("validate", help="校验 run 的结构、证据链和确认 Gate。")
    validate_parser.add_argument("--run", required=True, help="run 目录。")
    validate_parser.add_argument(
        "--strict-confirmations",
        action="store_true",
        help="要求两个确认 Gate 已完成且不存在 proposed 改写。",
    )
    validate_parser.set_defaults(handler=None)

    finalize_parser = subparsers.add_parser("finalize", help="校验、应用改写并导出 HTML、PDF 和日志。")
    finalize_parser.add_argument("--run", required=True, help="run 目录。")
    finalize_parser.set_defaults(handler=finalize_run)

    status_parser = subparsers.add_parser("status", help="查看 run 当前状态。")
    status_parser.add_argument("--run", required=True, help="run 目录。")
    status_parser.set_defaults(handler=status_run)

    promote_parser = subparsers.add_parser(
        "promote",
        help="将已确认的 run 事实同步到长期事实库。",
    )
    promote_parser.add_argument("--run", required=True, help="run 目录。")
    promote_parser.set_defaults(handler=promote_facts)
    return parser


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()
    try:
        if args.command == "validate":
            return 0 if print_validation(Path(args.run), args.strict_confirmations) else 1
        payload = args.handler(args)
        print(json.dumps(payload, ensure_ascii=False, indent=2))
        return 0
    except AgentError as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1
    except KeyboardInterrupt:
        print("ERROR: 操作已中断。", file=sys.stderr)
        return 130


if __name__ == "__main__":
    raise SystemExit(main())
